#!/usr/bin/env python3
"""Englische Etsy-Versionen aus den geprüften deutschen Excel-Vorlagen erzeugen.

Übersetzt Zelltexte, Formel-Literale, Blattnamen und Auswahllisten. Die Formeln selbst bleiben unverändert
→ die Rechnung ist dieselbe, die gegen tools/*/engine.js geprüft wurde. Währung neutral (#,##0).

    python3 tools/etsy/budget_xlsx.py && python3 tools/etsy/schulden_xlsx.py
    python3 tools/etsy/uebersetzen_en.py
    → content/etsy/en/budget-planner/files/Budget-Planner.xlsx
    → content/etsy/en/debt-payoff-planner/files/Debt-Payoff-Planner.xlsx
"""
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
GRAU, TINTE = "6B7280", "1F2937"

BLATT = {"Jahres-Tracker": "Yearly Tracker", "Anleitung": "How to use", "Listen": "Lists",
         "Monatsplan": "Monthly Plan", "Rechnung andere Methode": "Calc other method", "Rechnung ohne Extra": "Calc no extra"}
LITERAL = {"monatlich": "monthly", "pro Quartal": "quarterly", "halbjährlich": "semi-annual", "jährlich": "annual",
           "Muss": "Need", "Kann": "Want", "Sparen": "Save", "ja": "yes", "nicht erreichbar": "not reachable",
           "Lawine": "Avalanche", "Schneeball": "Snowball", "Rate deckt Zins nicht!": "Payment below interest!",
           "über 30 J.": "30+ yrs", "über 30 Jahre": "30+ years"}
TEXT = {
    # Budget
    "Budget-Plan Schweiz": "Budget Planner", "Einnahmen": "Income", "Ausgaben": "Expenses", "Bezeichnung": "Item",
    "Betrag": "Amount", "Takt": "Frequency", "Art": "Type", "Wohnen": "Housing", "pro Monat": "per month",
    "davon zurücklegen": "sinking fund / mo", "Nettolohn": "Take-home pay", "13. Monatslohn": "Annual bonus",
    "Miete inkl. Nebenkosten": "Rent / mortgage incl. utilities", "Krankenkasse Prämien": "Health insurance",
    "Franchise und Selbstbehalt": "Medical deductible", "Steuern (letzte Rechnung)": "Taxes not withheld (annual bill)",
    "Hausrat und Haftpflicht": "Home & liability insurance", "Radio- und TV-Gebühr": "Car registration & fees",
    "Handy und Internet": "Phone & internet", "ÖV und Mobilität": "Transport & fuel",
    "Lebensmittel und Haushalt": "Groceries & household", "Kleider": "Clothing", "Freizeit und Ausgang": "Going out & hobbies",
    "Ferien": "Vacation", "Geschenke": "Gifts & holidays", "Säule 3a": "Retirement savings",
    "Auswertung pro Monat": "Monthly summary", "Ausgaben (inkl. Sparen)": "Expenses (incl. saving)",
    "Übrig (+) / fehlt (−)": "Left over (+) / short (−)", "Dauerauftrag aufs Rückstellungskonto": "Monthly transfer to sinking funds",
    "Sparen (inkl. Überschuss)": "Save (incl. left over)", "Wohnkosten in % des Einkommens": "Housing as % of income",
    "Notreserve heute (auf dem Sparkonto)": "Emergency fund today", "Notreserve-Ziel (3 × Muss)": "Emergency fund goal (3 × Needs)",
    "Monate bis zum Ziel": "Months to goal", "Faustregel 50 / 30 / 20": "Rule of thumb 50 / 30 / 20",
    "Gelbe Felder ausfüllen, alles andere rechnet sich selbst. Takt = wie oft der Betrag anfällt.":
        "Fill in the yellow cells, everything else calculates itself. Frequency = how often you pay it.",
    "Tipp: Den Betrag „Dauerauftrag aufs Rückstellungskonto“ jeden Monat direkt nach dem Lohn auf ein separates Konto überweisen. Steuern, Franchise und Jahresrechnungen von dort bezahlen.":
        "Tip: move the „monthly transfer to sinking funds“ to a separate account right after payday. Pay all annual and irregular bills from there.",
    "Jahres-Tracker: was wirklich rausging": "Yearly tracker: what you actually spent",
    "Trag am Monatsende die tatsächlichen Beträge ein. Rot = mehr als geplant.": "Enter your actual spending at month end. Red = more than planned.",
    "Posten": "Item", "Plan / Monat": "Plan / month", "Ø Ist": "Avg actual", "Summe": "Total",
    "Jan": "Jan", "Feb": "Feb", "Mär": "Mar", "Apr": "Apr", "Mai": "May", "Jun": "Jun", "Jul": "Jul", "Aug": "Aug",
    "Sep": "Sep", "Okt": "Oct", "Nov": "Nov", "Dez": "Dec",
    # Schulden
    "Schulden-Plan Schweiz": "Debt Payoff Planner", "Schuld": "Debt", "Offen (CHF)": "Balance", "Zins % p. a.": "Interest % / yr",
    "Rate pro Monat": "Minimum payment", "Schuldenfrei (Monat)": "Paid off (month)", "Warnung": "Warning",
    "Extra pro Monat (CHF)": "Extra per month", "Methode": "Method", "Startmonat": "Start month",
    "Lawine = höchster Zins zuerst (günstigste) · Schneeball = kleinste Schuld zuerst":
        "Avalanche = highest interest first (cheapest) · Snowball = smallest balance first",
    "Gelbe Felder ausfüllen. Der Plan rechnet Monat für Monat, wann du schuldenfrei bist.":
        "Fill in the yellow cells. The plan works out, month by month, when you will be debt free.",
    "Kreditkarte": "Credit card", "Kleinkredit": "Personal loan", "Auto-Leasing": "Car loan", "Darlehen Familie": "Family loan",
    "Ergebnis": "Result", "Schuldenfrei nach (Monaten)": "Debt free after (months)", "Schuldenfrei im": "Debt free in",
    "Zinsen total (gewählte Methode)": "Total interest (chosen method)", "Zinsen mit der anderen Methode": "Interest with the other method",
    "Zinsen nur mit Mindestraten": "Interest with minimum payments only", "Ersparnis durch das Extra": "Saved by the extra payment",
    "Schuldenfrei nur mit Mindestraten (Monate)": "Debt free with minimums only (months)",
    "Tipp: Die Lawine kostet am wenigsten Zins. Der Schneeball bringt schneller ein erstes Erfolgserlebnis. Nimm die Methode, bei der du durchhältst.":
        "Tip: the avalanche costs the least interest. The snowball gives you a quick first win. Pick the method you will stick with.",
    "Extra pro Monat": "Extra per month", "Budget pro Monat": "Budget per month", "Start": "Start", "Zins": "Interest",
    "Mindestrate": "Minimum", "Rest": "Remaining", "Extra": "Extra", "Ende": "End", "Rang": "Rank", "Monat": "Month",
    "frei für Extra": "free for extra", "frei ab Monat": "paid off in month", "Zinsen total": "Total interest",
    "Schuldenfrei nach Monaten": "Debt free after months", "offen total": "total owed",
    "© aban news · abannews.com": "© aban news · abannews.com",
}
ANLEITUNG = {
    "budget": ["Budget Planner — how to use it", "",
               "1. Sheet „Budget“: enter your income. Put a yearly bonus on its own line with frequency „annual“.",
               "2. Enter every expense the way you pay it: rent monthly, insurance annually, car registration annually.",
               "   The „per month“ column converts everything to a monthly amount.",
               "3. Pick a type: Need (can't change short term), Want (can change), Save (retirement, savings, extra debt payments).",
               "4. Read the summary: left over / short, monthly transfer to sinking funds, Need/Want/Save, emergency fund.",
               "5. Sheet „Yearly Tracker“: at month end, enter what you actually spent. Red = more than 10 % over plan.", "",
               "Sinking funds: annual and irregular bills (insurance, taxes not withheld, car costs, holidays) break most budgets.",
               "Transfer the monthly amount from the summary to a separate account and pay those bills from there.", "",
               "Warning signs: left over is negative (red) or housing costs more than a third of your income (red).",
               "Works with any currency. Google Sheets: upload to Google Drive → Open with → Google Sheets.", "",
               "Planning tool, not financial advice. Personal use only. No sharing or reselling.", "© aban news · abannews.com"],
    "schulden": ["Debt Payoff Planner — how to use it", "",
                 "1. Sheet „Plan“: enter up to 8 debts: balance, yearly interest rate from your contract, current minimum payment.",
                 "2. Enter the extra amount you can pay each month. Even 50 a month shortens the payoff noticeably.",
                 "3. Pick a method: Avalanche (highest interest first, cheapest) or Snowball (smallest balance first).",
                 "4. Read the result: debt-free date, total interest, comparison with the other method and with minimums only.",
                 "5. Sheet „Monthly Plan“: month by month, what happens to each debt. Freed-up payments roll to the next debt.", "",
                 "Red in the „Warning“ column: the payment does not even cover the interest, so the debt grows. Fix that first.",
                 "Build an emergency fund of about three months of expenses first, then pay down debt.",
                 "Works with any currency. Google Sheets: upload to Google Drive → Open with → Google Sheets.", "",
                 "Model calculation without fees or penalty interest, not financial advice. Personal use only.", "© aban news · abannews.com"],
}


def formel(f):
    for alt, neu in BLATT.items():
        f = f.replace(f"'{alt}'!", f"'{neu}'!").replace(f"{alt}!", f"'{neu}'!" if " " in neu else f"{neu}!")
    return re.sub(r'"([^"]*)"', lambda m: '"' + LITERAL.get(m.group(1), m.group(1)) + '"', f)


def uebersetzen(quelle, ziel, art):
    wb = openpyxl.load_workbook(quelle)
    fehlend = set()
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    if v.startswith("="):
                        c.value = formel(v)
                    elif v in LITERAL:
                        c.value = LITERAL[v]
                    elif v in TEXT:
                        c.value = TEXT[v]
                    elif ws.title != "Anleitung" and v.strip():
                        fehlend.add(v)
                if c.number_format and "CHF" in c.number_format:
                    c.number_format = "#,##0;[Red]-#,##0" if "0.00" not in c.number_format else "#,##0.00"
        for dv in ws.data_validations.dataValidation:
            dv.formula1 = formel(dv.formula1) if dv.formula1.startswith("=") else re.sub(
                r"[^,\"]+", lambda m: LITERAL.get(m.group(0), m.group(0)), dv.formula1)
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                rule.formula = [formel(x) for x in rule.formula]
    h = wb["Anleitung"]
    h.delete_rows(1, h.max_row)
    for i, z in enumerate(ANLEITUNG[art], 1):
        c = h.cell(i, 1, z)
        c.font = Font(bold=True, size=18, color=TINTE) if i == 1 else Font(color=GRAU if z.startswith(("Planning", "Model", "©")) else TINTE)
    for alt, neu in BLATT.items():
        if alt in wb.sheetnames:
            wb[alt].title = neu
    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)
    print(f"→ {ziel.relative_to(ROOT)}" + (f"  ⚠ unübersetzt: {sorted(fehlend)}" if fehlend else "  (alle Texte übersetzt)"))
    return not fehlend


if __name__ == "__main__":
    ok = uebersetzen(ROOT / "content/etsy/budget-plan/dateien/Budget-Plan-Schweiz.xlsx",
                     ROOT / "content/etsy/en/budget-planner/files/Budget-Planner.xlsx", "budget")
    ok &= uebersetzen(ROOT / "content/etsy/schulden-plan/dateien/Schulden-Plan-Schweiz.xlsx",
                      ROOT / "content/etsy/en/debt-payoff-planner/files/Debt-Payoff-Planner.xlsx", "schulden")
    raise SystemExit(0 if ok else 1)
