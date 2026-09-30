#!/usr/bin/env python3
"""Hochzeits-Budget als Excel-Vorlage, Deutsch und Englisch. Alle Zahlen sind Formeln.

    python3 tools/etsy/hochzeit_xlsx.py
    → content/packs/de/hochzeits-budget/Hochzeits-Budget.xlsx          (im Kauf-Paket)
    → content/etsy/hochzeits-budget/dateien/Hochzeits-Budget.xlsx      (Etsy DE)
    → content/etsy/en/wedding-budget-planner/files/Wedding-Budget-Planner.xlsx (Etsy EN)

Rechenlogik wie tools/hochzeit/engine.js: Kosten = Betrag × Gäste (pro Gast) oder Betrag; Anzahlung am 1. des
Monats X Monate vorher (frühestens heute), Rest am Hochzeitstag; bereits Bezahltes deckt zuerst die Anzahlung.
Nötige Sparrate = grösster Wert von (bis dahin fällig − gespart) / Monate bis dahin.
"""
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
AKZ, CREME, LINIE, TINTE, GRAU = "B45309", "FEF3C7", "E7DFD2", "1F2937", "6B7280"
kopf, kopf_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=AKZ)
eingabe, summe_fill = PatternFill("solid", fgColor="FFF2B8"), PatternFill("solid", fgColor=CREME)
rand = Border(bottom=Side(style="thin", color=LINIE))
GELD, PZ, DAT = "#,##0;[Red]-#,##0", "0%", "DD.MM.YYYY"
ZEILEN, MONATE = 40, 36

T = {
 "de": dict(titel="Hochzeits-Budget-Plan", hinweis="Gelbe Felder ausfüllen, alles andere rechnet sich selbst.",
   budget="Budget", plan="Zahlungsplan", anl="Anleitung", ja="ja", nein="nein",
   eck=["Gäste", "Budget total", "Hochzeitsdatum", "Heute schon gespart", "Ihr spart pro Monat"],
   spalten=["Posten", "Betrag", "pro Gast", "Kosten", "Anzahlung %", "Monate vorher", "Anzahlung fällig", "Anzahlung", "schon bezahlt", "offen", "Rest am Hochzeitstag", "Anzahlung offen"],
   erg=["Kosten total", "pro Gast", "Übrig (+) / über Budget (−)", "Jeder zusätzliche Gast kostet", "Höchstens so viele Gäste im Budget", "Noch zu bezahlen", "Monate bis zur Hochzeit", "Nötige Sparrate pro Monat", "Mit eurer Sparrate: Lücke (−) / Reserve (+)"],
   warn="Sofort fällig ist mehr, als ihr gespart habt!",
   pspalten=["Monat", "fällig in diesem Monat", "fällig bis dann", "Monate ab heute", "angespart bis dann", "Lücke (−) / Reserve (+)", "nötige Rate bis hier"],
   phinweis="Rot = mit eurer Sparrate fehlt bis zu diesem Monat Geld. Anzahlungen am 1. des Monats, Rest im Hochzeitsmonat.",
   posten=[("Location / Saalmiete", 3000, 0, .5, 12), ("Essen", 120, 1, .3, 3), ("Getränke", 45, 1, 0, 0), ("Apéro", 25, 1, 0, 0),
     ("Hochzeitstorte", 10, 1, .3, 2), ("Fotografie", 3500, 0, .3, 9), ("Musik / DJ", 2000, 0, .3, 6), ("Kleid und Anzug", 3200, 0, .5, 8),
     ("Ringe", 2500, 0, 1, 3), ("Blumen und Deko", 1800, 0, .3, 4), ("Einladungen und Papeterie", 8, 1, 1, 5), ("Zivilstandsamt", 350, 0, 1, 2),
     ("Reserve für Unvorhergesehenes", 2000, 0, 0, 0)],
   anleitung=["Hochzeits-Budget-Plan — so geht's", "",
     "1. Blatt „Budget“: Gäste, Budget, Datum, schon Gespartes und eure monatliche Sparrate eintragen.",
     "2. Posten mit den Preisen aus euren Offerten. „pro Gast“ = ja, wenn der Preis pro Person gilt (Essen, Getränke, Torte).",
     "3. Anzahlung: wie viel Prozent der Anbieter vorher will und wie viele Monate vor der Hochzeit. Steht in der Offerte.",
     "4. Schon bezahlt: bereits überwiesene Beträge. Sie decken zuerst die Anzahlung, dann den Rest.",
     "5. Oben rechts: Kosten pro Gast, wie viele Gäste ins Budget passen und die Sparrate, die jede Zahlung deckt.",
     "6. Blatt „Zahlungsplan“: Monat für Monat, was fällig wird. Rot = mit eurer Sparrate fehlt dann Geld.", "",
     "Der grösste Hebel ist die Gästeliste: Pauschalen bleiben gleich, Essen und Getränke wachsen mit jedem Gast.",
     "Die Beispielbeträge sind Platzhalter, keine Durchschnittspreise. Funktioniert mit jeder Währung.",
     "Google Tabellen: Datei in Google Drive hochladen → Öffnen mit → Google Tabellen.", "",
     "Private Nutzung. Weitergabe und Weiterverkauf nicht gestattet.", "© aban news · abannews.com"]),
 "en": dict(titel="Wedding Budget Planner", hinweis="Fill in the yellow cells, everything else calculates itself.",
   budget="Budget", plan="Payment Plan", anl="How to use", ja="yes", nein="no",
   eck=["Guests", "Total budget", "Wedding date", "Saved so far", "You save per month"],
   spalten=["Item", "Amount", "per guest", "Cost", "Deposit %", "Months before", "Deposit due", "Deposit", "Already paid", "Open", "Balance on the day", "Deposit open"],
   erg=["Total cost", "per guest", "Left (+) / over budget (−)", "Each extra guest costs", "Max guests within budget", "Still to pay", "Months until the wedding", "Savings needed per month", "With your savings: gap (−) / buffer (+)"],
   warn="More is due right now than you have saved!",
   pspalten=["Month", "due this month", "due until then", "months from now", "saved by then", "gap (−) / buffer (+)", "rate needed up to here"],
   phinweis="Red = with your monthly savings you are short by this month. Deposits on the 1st, balance in the wedding month.",
   posten=[("Venue", 3000, 0, .5, 12), ("Food", 120, 1, .3, 3), ("Drinks", 45, 1, 0, 0), ("Cocktail hour", 25, 1, 0, 0),
     ("Wedding cake", 10, 1, .3, 2), ("Photographer", 3500, 0, .3, 9), ("Music / DJ", 2000, 0, .3, 6), ("Dress and suit", 3200, 0, .5, 8),
     ("Rings", 2500, 0, 1, 3), ("Flowers and decor", 1800, 0, .3, 4), ("Invitations and stationery", 8, 1, 1, 5), ("Marriage license / officiant", 350, 0, 1, 2),
     ("Buffer for surprises", 2000, 0, 0, 0)],
   anleitung=["Wedding Budget Planner — how to use it", "",
     "1. Sheet „Budget“: enter guests, budget, wedding date, what you have saved and your monthly savings.",
     "2. Items with the prices from your quotes. „per guest“ = yes if the price is per person (food, drinks, cake).",
     "3. Deposit: what percentage the vendor wants upfront and how many months before the wedding. It is in the quote.",
     "4. Already paid: money you have transferred. It covers the deposit first, then the balance.",
     "5. Top right: cost per guest, how many guests fit your budget and the monthly savings that cover every payment.",
     "6. Sheet „Payment Plan“: month by month, what is due. Red = with your savings you are short at that point.", "",
     "Your guest list is the biggest lever: flat fees stay the same, food and drinks grow with every guest.",
     "The sample amounts are placeholders, not average prices. Works with any currency.",
     "Google Sheets: upload to Google Drive → Open with → Google Sheets.", "",
     "Personal use only. No sharing or reselling.", "© aban news · abannews.com"]),
}


def bauen(sp, ziel):
    t = T[sp]
    wb = Workbook()
    b = wb.active
    b.title = t["budget"]
    b.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKL", [30, 11, 10, 12, 12, 12, 14, 11, 13, 11, 16, 13]):
        b.column_dimensions[col].width = w
    b["A1"], b["A2"] = t["titel"], t["hinweis"]
    b["A1"].font, b["A2"].font = Font(bold=True, size=18, color=TINTE), Font(color=GRAU, italic=True)
    # Eckdaten B4:B8
    werte = [80, 35000, "=DATE(YEAR(TODAY())+1,MONTH(TODAY()),15)", 6000, 1500]
    for i, (lab, w) in enumerate(zip(t["eck"], werte)):
        r = 4 + i
        b.cell(r, 1, lab).font = Font(bold=True, color=TINTE)
        c = b.cell(r, 2, w)
        c.fill, c.number_format = eingabe, (DAT if i == 2 else GELD)
    G, BUD, DAT_C, GESP, RATE = "$B$4", "$B$5", "$B$6", "$B$7", "$B$8"
    # Tabelle ab Zeile 11
    k = 11
    for c, v in enumerate(t["spalten"], 1):
        z = b.cell(k, c, v)
        z.font, z.fill, z.alignment = kopf, kopf_fill, Alignment(horizontal="left" if c == 1 else "right", wrap_text=True)
    b.column_dimensions["L"].hidden = True
    b.column_dimensions["M"].hidden = True
    dv = DataValidation(type="list", formula1=f'"{t["ja"]},{t["nein"]}"', allow_blank=True)
    b.add_data_validation(dv)
    von, bis = k + 1, k + ZEILEN
    for i in range(ZEILEN):
        r = von + i
        p = t["posten"][i] if i < len(t["posten"]) else None
        for c, v in zip((1, 2, 3, 5, 6, 9), (p[0], p[1], t["ja"] if p[2] else t["nein"], p[3], p[4], 0) if p else (None,) * 6):
            z = b.cell(r, c, v)
            z.fill = eingabe
        dv.add(f"C{r}")
        b[f"B{r}"].number_format = b[f"I{r}"].number_format = GELD
        b[f"E{r}"].number_format = PZ
        b[f"D{r}"] = f'=IF(C{r}="{t["ja"]}",N(B{r})*{G},N(B{r}))'
        b[f"H{r}"] = f"=D{r}*MIN(1,MAX(0,N(E{r})))"
        b[f"G{r}"] = f'=IF(H{r}-MIN(D{r},N(I{r}))>0.004,MAX(TODAY(),DATE(YEAR({DAT_C}),MONTH({DAT_C})-N(F{r}),1)),"")'
        b[f"J{r}"] = f"=D{r}-MIN(D{r},N(I{r}))"
        b[f"L{r}"] = f"=MAX(0,H{r}-MIN(D{r},N(I{r})))"
        b[f"K{r}"] = f"=J{r}-L{r}"
        b[f"M{r}"] = f'=IF(G{r}="","",YEAR(G{r})*12+MONTH(G{r}))'
        for c in "DHJKL":
            b[f"{c}{r}"].number_format = GELD
        b[f"G{r}"].number_format = DAT
        for c in range(1, 13):
            b.cell(r, c).border = rand
    # Ergebnis rechts oben (N/O)
    b.column_dimensions["N"].width, b.column_dimensions["O"].width = 38, 14
    rng = lambda c: f"{c}{von}:{c}{bis}"
    pg = f'SUMIF({rng("C")},"{t["ja"]}",{rng("B")})'
    formeln = [f"=SUM({rng('D')})", f"=IF({G}>0,O4/{G},0)", f"={BUD}-O4", f"={pg}",
               f'=IF({pg}>0,MAX(0,INT(({BUD}-(O4-{pg}*{G}))/{pg})),"–")', f"=SUM({rng('J')})",
               f"=MAX(0,(YEAR({DAT_C})-YEAR(TODAY()))*12+MONTH({DAT_C})-MONTH(TODAY()))",
               f"=MAX('{t['plan']}'!G5:G{4 + MONATE})",
               f"=MIN('{t['plan']}'!F5:F{4 + MONATE})"]
    for i, (lab, f) in enumerate(zip(t["erg"], formeln)):
        r = 4 + i
        b.cell(r, 14, lab).font = Font(bold=True, color=TINTE)
        z = b.cell(r, 15, f)
        z.fill, z.number_format, z.font = summe_fill, (GELD if i != 6 else "0"), Font(bold=True, color=TINTE)
    b["N14"] = f'=IF(\'{t["plan"]}\'!B5>{GESP},"{t["warn"]}","")'
    b["N14"].font = Font(bold=True, color="B91C1C")
    b.conditional_formatting.add("O6", CellIsRule(operator="lessThan", formula=["0"], font=Font(bold=True, color="B91C1C")))
    b.conditional_formatting.add("O12", CellIsRule(operator="lessThan", formula=["0"], font=Font(bold=True, color="B91C1C")))
    b.freeze_panes = "B12"

    # Zahlungsplan: Monat 0 = aktueller Monat
    z = wb.create_sheet(t["plan"])
    z.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [14, 20, 18, 16, 18, 22, 22]):
        z.column_dimensions[col].width = w
    z["A1"], z["A2"] = t["plan"], t["phinweis"]
    z["A1"].font, z["A2"].font = Font(bold=True, size=16, color=TINTE), Font(color=GRAU, italic=True)
    for c, v in enumerate(t["pspalten"], 1):
        h = z.cell(4, c, v)
        h.font, h.fill, h.alignment = kopf, kopf_fill, Alignment(horizontal="left" if c == 1 else "right", wrap_text=True)
    B = f"'{t['budget']}'!"
    for m in range(MONATE):
        r = 5 + m
        z[f"A{r}"] = f"=DATE(YEAR(TODAY()),MONTH(TODAY())+{m},1)"
        z[f"A{r}"].number_format = "MM.YYYY"
        # Monate nach dem Hochzeitsmonat: leer
        aktiv = f"A{r}<=DATE(YEAR({B}$B$6),MONTH({B}$B$6),1)"
        anz = f"SUMIF({B}$M${von}:$M${bis},YEAR(A{r})*12+MONTH(A{r}),{B}$L${von}:$L${bis})"
        rest = f"IF(AND(YEAR(A{r})=YEAR({B}$B$6),MONTH(A{r})=MONTH({B}$B$6)),SUM({B}$K${von}:$K${bis}),0)"
        z[f"B{r}"] = f'=IF({aktiv},{anz}+{rest},"")'
        z[f"C{r}"] = f'=IF({aktiv},SUM($B$5:B{r}),"")'
        z[f"D{r}"] = f'=IF({aktiv},{m},"")'
        z[f"E{r}"] = f'=IF({aktiv},{B}$B$7+{B}$B$8*{m},"")'
        z[f"F{r}"] = f'=IF({aktiv},E{r}-C{r},"")'
        z[f"G{r}"] = f'=IF({aktiv},IF({m}>0,MAX(0,(C{r}-{B}$B$7)/{m}),0),"")'
        for c in "BCEFG":
            z[f"{c}{r}"].number_format = GELD
        for c in range(1, 8):
            z.cell(r, c).border = rand
    z.conditional_formatting.add(f"A5:G{4 + MONATE}", FormulaRule(formula=['AND(ISNUMBER($F5),$F5<0)'], font=Font(bold=True, color="B91C1C")))
    z.freeze_panes = "A5"

    a = wb.create_sheet(t["anl"])
    a.column_dimensions["A"].width = 120
    for i, s in enumerate(t["anleitung"], 1):
        c = a.cell(i, 1, s)
        c.font = Font(bold=True, size=16, color=TINTE) if i == 1 else Font(color=GRAU if s.startswith(("©", "Private", "Personal")) else TINTE)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)
    print(f"→ {ziel.relative_to(ROOT)}")


if __name__ == "__main__":
    de = ROOT / "content/packs/de/hochzeits-budget/Hochzeits-Budget.xlsx"
    bauen("de", de)
    etsy = ROOT / "content/etsy/hochzeits-budget/dateien/Hochzeits-Budget.xlsx"
    etsy.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(de, etsy)
    print(f"→ {etsy.relative_to(ROOT)}")
    bauen("en", ROOT / "content/etsy/en/wedding-budget-planner/files/Wedding-Budget-Planner.xlsx")
